# -*- coding: utf-8 -*-
"""
Spoor 1 - Tevredenheid met het leven, toekomstperspectief en grip
=================================================================
Beantwoordt de subvragen van spoor 1 en zet ALLE cijfers in een Excel-bestand
(een tabblad per tabel/figuur) zodat de vormgever er direct een feitenblad van
kan maken.

Gebruik (Anaconda prompt, in deze map):
    python spoor1_analyse.py
of met een ander databestand:
    python spoor1_analyse.py "pad\\naar\\bestand.sav"

Output (in dezelfde map als dit script):
    spoor1_factsheet_data.xlsx   -> alle tabellen voor de vormgever
    spoor1_checks.txt            -> controle-log: gevonden kolommen, hercoderingen,
                                    frequenties. Even doorlezen na de eerste run!
    spoor1_open_antwoorden_INTERN.xlsx -> open antwoorden 'zorgen voor de toekomst'
                                    per groep, voor citaten. INTERN, eerst anonimiteit checken.

Alle percentages en gemiddelden zijn gewogen (kolom 'weegfactor').
"n" in de tabellen is altijd het ONGEWOGEN aantal respondenten.
"""
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pyreadstat
import statsmodels.api as sm

# ════════════════════════════════════════════════════════════════════════════
# INSTELLINGEN
# ════════════════════════════════════════════════════════════════════════════
DATA_PATH = r"O:\Projecten lopend\BW & kansengelijkheid\Data\Werkbestand BW en kansengelijkheid.sav"
OUT_DIR = Path(__file__).resolve().parent
OUT_XLSX = OUT_DIR / "spoor1_factsheet_data.xlsx"
OUT_LOG = OUT_DIR / "spoor1_checks.txt"

W = "weegfactor"
TEVREDEN_GRENS = 7      # rapportcijfer >= 7  -> tevreden
ONVOLDOENDE_GRENS = 6   # rapportcijfer <  6  -> domein 'niet op orde' (onvoldoende)
LEEFTIJD_BINS = [17, 34, 49, 64, 200]
LEEFTIJD_LABELS = ["18-34 jaar", "35-49 jaar", "50-64 jaar", "65 jaar en ouder"]

GROEPEN = [
    "Tevreden & hoopvol/stabiel",
    "Tevreden & pessimistisch",
    "Ontevreden & hoopvol/stabiel",
    "Ontevreden & pessimistisch",
]
TOTAAL = "Totaal"

# Korte, leesbare domeinnamen voor de vormgever (volgorde = volgorde in tabellen)
DOMEIN_KORT = [
    ("leven in het algemeen", "Leven in het algemeen"),
    ("financi", "Financiële situatie"),
    ("lichamelijke gezondheid", "Lichamelijke gezondheid"),
    ("psychische gezondheid", "Psychische gezondheid"),
    ("werk en/of andere", "Werk/dagelijkse bezigheden"),
    ("vrije tijd", "Vrije tijd"),
    ("opleidingskansen", "Opleidingskansen"),
    ("woning of woonruimte", "Woning"),
    ("directe woonomgeving", "Buurt/wijk"),
    ("sociale contacten", "Sociale contacten"),
    ("met elkaar omgaan", "Omgang in de buurt"),
    ("veiligheid", "Veiligheid in de buurt"),
    ("groen en de natuur", "Groen en natuur"),
    ("milieukwaliteit", "Milieukwaliteit"),
    ("bereikbaarheid", "Bereikbaarheid voorzieningen"),
    ("culturele aanbod", "Cultureel aanbod"),
]
LEVEN = "Leven in het algemeen"

TOEKOMST_MAP = {"veel slechter": 1, "een beetje slechter": 2, "hetzelfde": 3,
                "een beetje beter": 4, "veel beter": 5}
TOEKOMST_LABELS = ["Veel slechter", "Een beetje slechter", "Hetzelfde",
                   "Een beetje beter", "Veel beter"]
INVLOED_MAP = {"geen invloed": 0, "een beetje invloed": 1, "veel invloed": 2}

_LOG = []


def log(*args):
    msg = " ".join(str(a) for a in args)
    print(msg)
    _LOG.append(msg)


# ════════════════════════════════════════════════════════════════════════════
# 1. DATA INLEZEN
# ════════════════════════════════════════════════════════════════════════════
VALUE_ORDER = {}   # kolom -> volgorde van de SPSS-waardelabels (voor nette tabellen)


def load_and_label(path):
    df, meta = pyreadstat.read_sav(path)
    label_map = {c: l for c, l in zip(meta.column_names, meta.column_labels) if l}
    df.rename(columns=label_map, inplace=True)
    for col_original, col_renamed in label_map.items():
        if col_original in meta.variable_value_labels:
            value_map = meta.variable_value_labels[col_original]
            df[col_renamed] = df[col_renamed].map(value_map).fillna(df[col_renamed])
            VALUE_ORDER[col_renamed] = list(dict.fromkeys(value_map.values()))
    for col in meta.column_names:  # ongelabelde kolommen met waardelabels
        if col not in label_map and col in meta.variable_value_labels:
            value_map = meta.variable_value_labels[col]
            df[col] = df[col].map(value_map).fillna(df[col])
            VALUE_ORDER[col] = list(dict.fromkeys(value_map.values()))
    # dubbele kolomnamen na hernoemen: eerste houden, rest suffix geven
    cols = pd.Series(df.columns)
    for dup in cols[cols.duplicated()].unique():
        for i, idx in enumerate(cols[cols == dup].index):
            if i > 0:
                cols[idx] = f"{dup}__{i + 1}"
    df.columns = cols
    return df


def load_data(path):
    df = load_and_label(path)
    n0 = len(df)
    if "Reactiestatus" in df.columns:
        df = df[df["Reactiestatus"] != "Niet geantwoord"].copy()
    log(f"Ingelezen: {n0} rijen, na filter Reactiestatus: {len(df)} rijen, {df.shape[1]} kolommen")
    return df


# ════════════════════════════════════════════════════════════════════════════
# 2. HULPFUNCTIES
# ════════════════════════════════════════════════════════════════════════════
def clean_label(c):
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", str(c))).strip()


def subject(c):
    return clean_label(c).split(" - ")[-1].strip()


def find_cols(df, *needles, exclude=()):
    out = []
    for c in df.columns:
        lc = clean_label(c).lower()
        if all(n.lower() in lc for n in needles) and not any(e.lower() in lc for e in exclude):
            out.append(c)
    return out


def first_col(df, *candidates):
    """Eerste kolom waarvan de naam (of label) een van de kandidaten bevat."""
    for cand in candidates:
        if cand in df.columns:
            return cand
    for cand in candidates:
        hits = find_cols(df, cand)
        if hits:
            return hits[0]
    return None


def kort(s):
    ls = s.lower()
    for key, name in DOMEIN_KORT:
        if key in ls:
            return name
    return s[:45]


def norm_str(s):
    return s.astype("string").str.strip().str.lower()


def to_score(s):
    """Rapportcijfer 1-10 numeriek maken (ook als de uitersten een label hebben)."""
    if pd.api.types.is_numeric_dtype(s):
        return s.astype(float)
    num = pd.to_numeric(s, errors="coerce")
    lead = pd.to_numeric(s.astype("string").str.extract(r"^\s*(\d{1,2})")[0], errors="coerce")
    return num.fillna(lead).astype(float)


def to_toekomst(s):
    if pd.api.types.is_numeric_dtype(s):
        log(f"  LET OP: toekomstkolom '{subject(s.name)}' is al numeriek, onveranderd gebruikt")
        return s.astype(float)
    return norm_str(s).map(TOEKOMST_MAP).astype(float)


def to_invloed(s):
    return norm_str(s).map(INVLOED_MAP).astype(float)


def to_likert(s):
    """(helemaal) oneens 1 ... (helemaal) eens 5; weet niet -> NaN."""
    if pd.api.types.is_numeric_dtype(s):
        return s.astype(float)

    def m(x):
        if x is pd.NA or x is None:
            return np.nan
        x = str(x).lower().replace("mee ", "")   # 'helemaal mee eens' -> 'helemaal eens'
        if "weet" in x or "wil ik niet" in x:
            return np.nan
        if "neutraal" in x or ("niet eens" in x and "niet oneens" in x):
            return 3
        if "helemaal oneens" in x or "zeer oneens" in x:
            return 1
        if "helemaal eens" in x or "zeer eens" in x:
            return 5
        if "oneens" in x:
            return 2
        if "eens" in x:
            return 4
        return np.nan
    return s.map(m).astype(float)


def to_janee(s):
    """Ja -> 1, Nee / Weet ik niet -> 0 (geen zekere steun), leeg -> NaN."""
    if pd.api.types.is_numeric_dtype(s):
        return (s > 0).astype(float).where(s.notna())
    x = norm_str(s)
    out = pd.Series(np.nan, index=s.index)
    out[x.str.startswith("ja", na=False)] = 1
    out[x.str.startswith("nee", na=False) | x.str.contains("weet", na=False)] = 0
    return out


def binary_block(df, cols):
    """Meerkeuze-aanvinkvragen -> 0/1. Wie geen enkele optie van het blok heeft ingevuld -> NaN."""
    out = pd.DataFrame(index=df.index)
    for c in cols:
        s = df[c]
        if pd.api.types.is_numeric_dtype(s):
            out[c] = (s > 0).astype(float).where(s.notna())
        else:
            x = norm_str(s)
            v = pd.Series(np.nan, index=s.index)
            v[x.isin(["0", "0.0", "nee", "niet geselecteerd", "not selected", "false"])] = 0
            v[x.notna() & v.isna()] = 1
            out[c] = v
    answered = out.notna().any(axis=1)
    out = out.apply(lambda col: col.fillna(0).where(answered))
    return out


def wmean(x, w):
    m = x.notna() & w.notna()
    if m.sum() == 0:
        return np.nan
    return float(np.average(x[m].astype(float), weights=w[m].astype(float)))


def wpct(x, w):
    v = wmean(x, w)
    return np.nan if pd.isna(v) else 100 * v


def wcorr(x, y, w):
    m = x.notna() & y.notna() & w.notna()
    if m.sum() < 10:
        return np.nan, int(m.sum())
    x, y, w = x[m].astype(float), y[m].astype(float), w[m].astype(float)
    mx, my = np.average(x, weights=w), np.average(y, weights=w)
    cov = np.average((x - mx) * (y - my), weights=w)
    r = cov / np.sqrt(np.average((x - mx) ** 2, weights=w) * np.average((y - my) ** 2, weights=w))
    return float(r), int(m.sum())


def wdist(cat, w, order=None):
    """Gewogen % per categorie."""
    m = cat.notna() & w.notna()
    tot = w[m].sum()
    cats = order if order is not None else sorted(cat[m].unique(), key=str)
    return {c: (100 * w[m & (cat == c)].sum() / tot if tot > 0 else np.nan) for c in cats}


def z(s):
    s = s.astype(float)
    return (s - s.mean()) / s.std()


def ordered_cats(df, col):
    present = [v for v in df[col].dropna().unique()]
    order = [v for v in VALUE_ORDER.get(col, []) if v in present]
    rest = sorted([v for v in present if v not in order], key=str)
    return order + rest


# ════════════════════════════════════════════════════════════════════════════
# 3. VARIABELEN OPBOUWEN
# ════════════════════════════════════════════════════════════════════════════
def build_variables(df):
    V = {}  # alle afgeleide variabelen + kolomlijsten

    # weegfactor
    if W not in df.columns:
        raise KeyError(f"Weegfactor-kolom '{W}' niet gevonden")
    df[W] = pd.to_numeric(df[W], errors="coerce")
    log(f"Weegfactor: {df[W].notna().sum()} geldig, gemiddelde {df[W].mean():.3f}")

    # ── tevredenheid nu & toekomst per domein ──────────────────────────────
    tev_cols = find_cols(df, "hoe tevreden u bent met")
    toe_cols = [c for c in find_cols(df, "komende 5 jaar") if clean_label(c).lower().startswith("hoe denkt u")]
    log(f"\nTevredenheidskolommen: {len(tev_cols)}  |  Toekomstkolommen: {len(toe_cols)}")

    tev = pd.DataFrame(index=df.index)
    toe = pd.DataFrame(index=df.index)
    for c in tev_cols:
        tev[kort(subject(c))] = to_score(df[c])
    for c in toe_cols:
        k = kort(subject(c))
        log(f"  toekomst '{k}': waarden {sorted(df[c].dropna().astype(str).unique())[:8]}")
        toe[k] = to_toekomst(df[c])
    domeinen = [k for _, k in DOMEIN_KORT if k in tev.columns]
    V["tev"], V["toe"], V["domeinen"] = tev, toe, domeinen
    V["domeinen_excl_leven"] = [d for d in domeinen if d != LEVEN]
    if LEVEN not in tev.columns or LEVEN not in toe.columns:
        raise KeyError("'Uw leven in het algemeen' niet gevonden in tevredenheid/toekomst")

    # ── ervaren invloed (grip) ──────────────────────────────────────────────
    inv = pd.DataFrame(index=df.index)
    V["inv_raw"] = {}
    for c in find_cols(df, "hoeveel invloed"):
        k = kort(subject(c))
        inv[k] = to_invloed(df[c])
        V["inv_raw"][k] = c
        log(f"  invloed '{k}': {df[c].value_counts(dropna=False).to_dict()}")
    V["inv"] = inv

    # ── stellingen over grip/kansen ─────────────────────────────────────────
    stel = pd.DataFrame(index=df.index)
    for c in find_cols(df, "in hoeverre u het eens bent"):
        k = subject(c)
        stel[k] = to_likert(df[c])
        log(f"  stelling '{k[:60]}': {df[c].value_counts(dropna=False).to_dict()}")
    V["stel"] = stel

    def _stel(key):
        return next((c for c in stel.columns if key in c.lower()), None)
    V["grip_stelling"] = _stel("invloed heb op hoe mijn leven")
    V["stel_kort"] = [(k, c) for k, c in [
        ("Zelf veel invloed op hoe leven verloopt", V["grip_stelling"]),
        ("Neemt zelf belangrijke beslissingen voor toekomst", _stel("zelf beslissingen")),
        ("Zelf minstens evenveel kansen gekregen", _stel("evenveel kansen gekregen")),
        ("In Nederland heeft iedereen evenveel kansen", _stel("iedereen evenveel kansen")),
    ] if c]

    # ── kernvariabelen ──────────────────────────────────────────────────────
    df["_tev_leven"] = tev[LEVEN]
    df["_toe_leven"] = toe[LEVEN]
    df["_tevreden"] = (df["_tev_leven"] >= TEVREDEN_GRENS).astype(float).where(df["_tev_leven"].notna())
    df["_onvoldoende_leven"] = (df["_tev_leven"] < ONVOLDOENDE_GRENS).astype(float).where(df["_tev_leven"].notna())
    df["_pessimistisch"] = (df["_toe_leven"] < 3).astype(float).where(df["_toe_leven"].notna())
    df["_optimistisch"] = (df["_toe_leven"] > 3).astype(float).where(df["_toe_leven"].notna())
    df["_inv_leven"] = inv[LEVEN] if LEVEN in inv.columns else np.nan
    df["_veel_invloed"] = (df["_inv_leven"] == 2).astype(float).where(df["_inv_leven"].notna())

    grp = pd.Series(pd.NA, index=df.index, dtype="object")
    t, p = df["_tevreden"], df["_pessimistisch"]
    grp[(t == 1) & (p == 0)] = GROEPEN[0]
    grp[(t == 1) & (p == 1)] = GROEPEN[1]
    grp[(t == 0) & (p == 0)] = GROEPEN[2]
    grp[(t == 0) & (p == 1)] = GROEPEN[3]
    df["_groep"] = grp
    log("\nIndeling groepen (ongewogen n):")
    log(df["_groep"].value_counts(dropna=False).to_string())
    q_lev = first_col(df, "Q_levenalgemeen")
    if q_lev:
        log("\nControle: eigen indeling vs bestaande variabele Q_levenalgemeen")
        log(pd.crosstab(df["_groep"].fillna("(leeg)"), df[q_lev].fillna("(leeg)")).to_string())

    # domein onvoldoende-dummies & aantal
    onv = pd.DataFrame(index=df.index)
    for d in V["domeinen_excl_leven"]:
        onv[d] = (tev[d] < ONVOLDOENDE_GRENS).astype(float).where(tev[d].notna())
    V["onv"] = onv
    df["_n_onvoldoende"] = onv.sum(axis=1, min_count=1)

    # ── bestaans(on)zekerheid & andere onzekerheden ─────────────────────────
    excl = ("geen van bovenstaande", "wil ik niet zeggen")
    fin_cols = find_cols(df, "situaties zijn op u van toepassing", exclude=excl)
    ovr_cols = find_cols(df, "situaties is op u van toepassing", exclude=excl)
    fin = binary_block(df, fin_cols)
    ovr = binary_block(df, ovr_cols)
    fin.columns = [subject(c) for c in fin_cols]
    ovr.columns = [subject(c) for c in ovr_cols]
    V["onz_fin"], V["onz_ovr"] = fin, ovr
    df["_n_onz_fin"] = fin.sum(axis=1, min_count=1)
    df["_n_onz_tot"] = pd.concat([fin, ovr], axis=1).sum(axis=1, min_count=1)
    log(f"\nOnzekerheidsitems: {len(fin_cols)} financieel/bestaans, {len(ovr_cols)} overig")
    for c in fin_cols[:2] + ovr_cols[:1]:
        log(f"  voorbeeld '{subject(c)[:50]}': {df[c].value_counts(dropna=False).to_dict()}")

    # ── sociaal netwerk / weten waar hulp ───────────────────────────────────
    net = pd.DataFrame(index=df.index)
    for c in find_cols(df, "zou kunnen helpen met") + find_cols(df, "terechtkunt"):
        net[subject(c)] = to_janee(df[c])
    hulp = pd.DataFrame(index=df.index)
    for c in find_cols(df, "welke organisaties ik terecht kan"):
        hulp[subject(c)] = to_janee(df[c])
    if len(net.columns):
        log(f"  netwerk voorbeeld: {df[find_cols(df, 'zou kunnen helpen met')[0]].value_counts(dropna=False).to_dict()}")
    if len(hulp.columns):
        log(f"  weet-waar-hulp voorbeeld: {df[find_cols(df, 'welke organisaties ik terecht kan')[0]].value_counts(dropna=False).to_dict()}")
    V["net"], V["hulp"] = net, hulp

    # ── achtergrondkenmerken ────────────────────────────────────────────────
    df["_leeftijd"] = pd.to_numeric(df.get("Leeftijd"), errors="coerce")
    df["Leeftijdsgroep"] = pd.cut(df["_leeftijd"], LEEFTIJD_BINS, labels=LEEFTIJD_LABELS).astype("object")
    VALUE_ORDER["Leeftijdsgroep"] = LEEFTIJD_LABELS

    gesl = first_col(df, "Geslacht")
    if gesl:
        g = df[gesl].astype("string").str.strip()
        g = g.replace({"Een man": "Man", "Een vrouw": "Vrouw", "een man": "Man", "een vrouw": "Vrouw"})
        df["Geslacht (geharmoniseerd)"] = g.astype("object")
        VALUE_ORDER["Geslacht (geharmoniseerd)"] = ["Man", "Vrouw", "Non-binair", "Wil ik niet zeggen"]

    achtergrond = [  # (weergavenaam, kolom)
        ("Leeftijdsgroep", "Leeftijdsgroep"),
        ("Geslacht", "Geslacht (geharmoniseerd)" if gesl else None),
        ("Opleidingsniveau", first_col(df, "opleiding3", "Opleiding (3 niveaus)")),
        ("Inkomen huishouden", first_col(df, "Inkomencat", "Inkomen (categorie)")),
        ("Koop/huur", first_col(df, "Woont u in een koopwoning of huurwoning")),
        ("Huishouden", first_col(df, "Hoe kunt u uw huishouden het beste omschrijven")),
        ("Belangrijkste inkomensbron", first_col(df, "Uit welke bron komt het grootste deel")),
        ("Provincie", first_col(df, "Provincie", "regioprovincie")),
    ]
    V["achtergrond"] = [(n, c) for n, c in achtergrond if c]
    V["dagbesteding"] = binary_block(df, find_cols(df, "in welke situatie(s) zit u", exclude=("wil ik niet zeggen",)))
    V["dagbesteding"].columns = [subject(c) for c in V["dagbesteding"].columns]

    verleden = [
        ("Ouders konden goed rondkomen (jeugd)", first_col(df, "Konden uw ouders of verzorgers over het algemeen goed rondkomen")),
        ("Ouders hadden voldoende tijd/aandacht (jeugd)", first_col(df, "Hadden uw ouders of verzorgers voldoende tijd en aandacht")),
        ("Kon als volwassene (25+) goed rondkomen", first_col(df, "Kon u als volwassene (vanaf 25 jaar) meestal goed rondkomen")),
        ("Moeite met rondkomen afgelopen 12 maanden", first_col(df, "afgelopen 12 maanden moeite gehad om van het inkomen")),
    ]
    V["verleden"] = [(n, c) for n, c in verleden if c]

    log("\nAchtergrond gevonden:", [(n, c[:40]) for n, c in V["achtergrond"]])
    log("Verleden gevonden:", [n for n, _ in V["verleden"]])
    for n, c in V["achtergrond"] + V["verleden"]:
        log(f"  {n}: {df[c].value_counts(dropna=False).head(12).to_dict()}")
    missing = [n for n, c in achtergrond + verleden if not c]
    if missing:
        log("LET OP, niet gevonden:", missing)

    V["opl_col"] = dict(V["achtergrond"]).get("Opleidingsniveau")
    V["ink_col"] = dict(V["achtergrond"]).get("Inkomen huishouden")
    return V


def controles(df, V):
    """Controlevariabelen voor regressies: leeftijd, geslacht, opleiding, inkomen."""
    X = pd.DataFrame(index=df.index)
    X["Leeftijd"] = df["_leeftijd"]
    if "Geslacht (geharmoniseerd)" in df.columns:
        X["Vrouw"] = df["Geslacht (geharmoniseerd)"].map({"Vrouw": 1.0, "Man": 0.0})
    for col, prefix in [(V["opl_col"], "Opleiding"), (V["ink_col"], "Inkomen")]:
        if not col:
            continue
        s = df[col].astype("object")
        s = s.where(~s.astype(str).str.lower().str.contains("weet|wil ik niet|onbekend", na=False))
        if s.notna().sum() == 0:
            continue
        ref = s.value_counts().idxmax()
        d = pd.get_dummies(s, prefix=prefix, dtype=float)
        d = d.drop(columns=[f"{prefix}_{ref}"], errors="ignore")
        d[s.isna()] = np.nan
        X = pd.concat([X, d], axis=1)
    return X


def fit_wls(y, X, w):
    d = pd.concat([y.rename("_y"), X, w.rename("_w")], axis=1).dropna()
    model = sm.WLS(d["_y"], sm.add_constant(d[X.columns], has_constant="add"), weights=d["_w"]).fit()
    return model, len(d)


def fit_logit(y, X, w):
    d = pd.concat([y.rename("_y"), X, w.rename("_w")], axis=1).dropna()
    fw = d["_w"] / d["_w"].mean()
    model = sm.GLM(d["_y"], sm.add_constant(d[X.columns], has_constant="add"),
                   family=sm.families.Binomial(), freq_weights=fw).fit()
    return model, len(d)


# ════════════════════════════════════════════════════════════════════════════
# 4. TABELLEN
# ════════════════════════════════════════════════════════════════════════════
SHEETS = []


def add_sheet(name, titel, vraag, df, toelichting="", fmt=None, visual=""):
    SHEETS.append(dict(name=name[:31], titel=titel, vraag=vraag, df=df.reset_index(drop=True),
                       toelichting=toelichting, fmt=fmt or {}, visual=visual))


def groep_cols():
    return GROEPEN + [TOTAAL]


def per_groep(df, fn):
    """fn(mask) -> waarde; geeft dict per groep + totaal (totaal = alle ingedeelde respondenten)."""
    ingedeeld = df["_groep"].notna()
    out = {g: fn(df["_groep"] == g) for g in GROEPEN}
    out[TOTAAL] = fn(ingedeeld)
    return out


def n_row(df, label="n (ongewogen)", **extra):
    row = dict(extra)
    row.update({g: int((df["_groep"] == g).sum()) for g in GROEPEN})
    row[TOTAAL] = int(df["_groep"].notna().sum())
    return row


# ---------------------------------------------------------------- Subvraag 1
def tabellen_vraag1(df, V):
    w = df[W]
    VR = "1. Hangt (on)tevredenheid met het leven samen met de blik op de toekomst (en met grip)?"

    # 1a kerncijfers
    r_tt, n_tt = wcorr(df["_tev_leven"], df["_toe_leven"], w)
    r_ti, _ = wcorr(df["_tev_leven"], df["_inv_leven"], w)
    rows = [
        ("Aantal respondenten (ongewogen)", len(df), "aantal", ""),
        ("Gemiddelde tevredenheid met het leven", wmean(df["_tev_leven"], w), "rapportcijfer 1-10", ""),
        (f"Tevreden met het leven (cijfer {TEVREDEN_GRENS}+)", wpct(df["_tevreden"], w), "%", ""),
        (f"Onvoldoende voor het leven (cijfer <{ONVOLDOENDE_GRENS})", wpct(df["_onvoldoende_leven"], w), "%", ""),
        ("Verwacht dat het leven over 5 jaar beter is", wpct(df["_optimistisch"], w), "%", ""),
        ("Verwacht dat het leven over 5 jaar hetzelfde is", wpct((df["_toe_leven"] == 3).astype(float).where(df["_toe_leven"].notna()), w), "%", ""),
        ("Verwacht dat het leven over 5 jaar slechter is (pessimistisch)", wpct(df["_pessimistisch"], w), "%", ""),
        ("Ervaart veel invloed op eigen leven", wpct(df["_veel_invloed"], w), "%", ""),
        ("Ervaart geen invloed op eigen leven", wpct((df["_inv_leven"] == 0).astype(float).where(df["_inv_leven"].notna()), w), "%", ""),
        ("Samenhang tevredenheid nu en toekomstverwachting (r)", r_tt, "correlatie -1 tot 1",
         "0,1 = zwak, 0,3 = matig, 0,5 = sterk"),
        ("Samenhang tevredenheid nu en ervaren invloed (r)", r_ti, "correlatie -1 tot 1", ""),
    ]
    for g in GROEPEN:
        rows.append((f"Groep: {g}", wpct((df["_groep"] == g).astype(float).where(df["_groep"].notna()), w), "%",
                     f"n = {(df['_groep'] == g).sum()}"))
    add_sheet("1a_Kerncijfers", "Kerncijfers tevredenheid, toekomstperspectief en grip", VR,
              pd.DataFrame(rows, columns=["Indicator", "Waarde", "Eenheid", "Opmerking"]),
              "Gewogen. Basis voor de 'grote getallen' op het feitenblad.",
              fmt={"Waarde": "0.0"}, visual="Kerncijfers / iconen")

    # 1b toekomstverwachting per tevredenheidsniveau
    band = pd.cut(df["_tev_leven"], [0, 5, 6, 7, 8, 10], labels=["1-5", "6", "7", "8", "9-10"]).astype("object")
    toe_lab = df["_toe_leven"].map(dict(enumerate(TOEKOMST_LABELS, start=1)))
    rows = []
    for b in ["1-5", "6", "7", "8", "9-10"]:
        m = band == b
        d = wdist(toe_lab[m], w[m], TOEKOMST_LABELS)
        rows.append({"Tevredenheid met leven (cijfer)": b, **d,
                     "% pessimistisch (slechter)": d["Veel slechter"] + d["Een beetje slechter"],
                     "% optimistisch (beter)": d["Een beetje beter"] + d["Veel beter"],
                     "n": int((m & toe_lab.notna()).sum())})
    add_sheet("1b_Toekomst_per_tevredenheid", "Toekomstverwachting naar tevredenheid met het leven nu",
              VR, pd.DataFrame(rows),
              "Rij-percentages (tellen per rij op tot 100%). 'Weet ik niet' bij toekomst niet meegeteld.",
              visual="100%-gestapelde staaf per cijfergroep")

    # 1c tevredenheid per toekomstverwachting en per invloed
    rows = []
    for i, lab in enumerate(TOEKOMST_LABELS, start=1):
        m = df["_toe_leven"] == i
        rows.append({"Kenmerk": "Verwachting leven over 5 jaar", "Categorie": lab,
                     "Gem. tevredenheid leven (1-10)": wmean(df.loc[m, "_tev_leven"], w[m]),
                     f"% tevreden ({TEVREDEN_GRENS}+)": wpct(df.loc[m, "_tevreden"], w[m]),
                     "% van alle respondenten": wpct(m.astype(float).where(df["_toe_leven"].notna()), w),
                     "n": int(m.sum())})
    for val, lab in [(0, "Geen invloed"), (1, "Een beetje invloed"), (2, "Veel invloed")]:
        m = df["_inv_leven"] == val
        rows.append({"Kenmerk": "Ervaren invloed op eigen leven", "Categorie": lab,
                     "Gem. tevredenheid leven (1-10)": wmean(df.loc[m, "_tev_leven"], w[m]),
                     f"% tevreden ({TEVREDEN_GRENS}+)": wpct(df.loc[m, "_tevreden"], w[m]),
                     "% van alle respondenten": wpct(m.astype(float).where(df["_inv_leven"].notna()), w),
                     "n": int(m.sum())})
        rows[-1]["% pessimistisch over leven"] = wpct(df.loc[m, "_pessimistisch"], w[m])
    for i in range(len(TOEKOMST_LABELS)):
        rows[i]["% pessimistisch over leven"] = np.nan
    add_sheet("1c_Tevredenheid_per_toekomst", "Tevredenheid met het leven naar toekomstverwachting en ervaren invloed",
              VR, pd.DataFrame(rows), "Gewogen gemiddelden/percentages.",
              fmt={"Gem. tevredenheid leven (1-10)": "0.0"}, visual="Staafdiagram gemiddeld cijfer")

    # 1d samenhang per domein (nu vs toekomst, nu vs invloed)
    rows = []
    for d in V["domeinen"]:
        r1, n1 = wcorr(V["tev"][d], V["toe"][d], w) if d in V["toe"].columns else (np.nan, 0)
        r2, _ = wcorr(V["tev"][d], V["inv"][d], w) if d in V["inv"].columns else (np.nan, 0)
        r3, _ = wcorr(V["tev"][d], df["_tev_leven"], w) if d != LEVEN else (np.nan, 0)
        r4, _ = wcorr(V["toe"][d], df["_toe_leven"], w) if (d != LEVEN and d in V["toe"].columns) else (np.nan, 0)
        rows.append({"Domein": d,
                     "r tevredenheid nu x verwachting toekomst (zelfde domein)": r1,
                     "r tevredenheid nu x ervaren invloed (zelfde domein)": r2,
                     "r domein-tevredenheid x tevredenheid leven": r3,
                     "r domein-verwachting x verwachting leven": r4,
                     "n": n1})
    rfmt = {c: "0.00" for c in rows[0] if c.startswith("r ")}
    add_sheet("1d_Samenhang_per_domein", "Samenhang (correlatie) tussen tevredenheid nu, toekomst en invloed per domein",
              VR, pd.DataFrame(rows),
              "Gewogen Pearson-correlaties. 0,1 = zwak, 0,3 = matig, 0,5 = sterk. Leeg = vraag niet gesteld voor dit domein.",
              fmt=rfmt, visual="Gesorteerde staaf of heatmap")

    # 1e regressie
    C = controles(df, V)
    preds = {"Verwachting leven over 5 jaar (1-5)": df["_toe_leven"],
             "Ervaren invloed op leven (0-2)": df["_inv_leven"]}
    if V["grip_stelling"]:
        preds["Stelling: 'ik heb veel invloed op hoe mijn leven verloopt' (1-5)"] = V["stel"][V["grip_stelling"]]
    base = pd.concat([df["_tev_leven"].rename("_y")] + [s.rename(k) for k, s in preds.items()] + [C, w], axis=1).dropna()
    idx = base.index
    y = z(df.loc[idx, "_tev_leven"])
    Xp = pd.DataFrame({k: z(s.loc[idx]) for k, s in preds.items()})
    Cz = C.loc[idx].copy()
    Cz["Leeftijd"] = z(Cz["Leeftijd"])
    modellen = [
        ("Model 1: alleen achtergrond", []),
        ("Model 2: + toekomstverwachting", [list(preds)[0]]),
        ("Model 3: + invloed/grip", list(preds)),
    ]
    res = {}
    for naam, p in modellen:
        X = pd.concat([Cz, Xp[p]], axis=1)
        mdl, n = fit_wls(y, X, w.loc[idx])
        res[naam] = (mdl, n)
    rows = []
    for k in list(preds) + ["Leeftijd"]:
        row = {"Voorspeller": k}
        for naam, _ in modellen:
            mdl = res[naam][0]
            row[f"{naam} - beta"] = mdl.params.get(k, np.nan)
            row[f"{naam} - p"] = mdl.pvalues.get(k, np.nan)
        rows.append(row)
    for stat in ["R² (verklaarde variantie)", "n"]:
        row = {"Voorspeller": stat}
        for naam, _ in modellen:
            mdl, n = res[naam]
            row[f"{naam} - beta"] = mdl.rsquared if stat.startswith("R") else n
            row[f"{naam} - p"] = np.nan
        rows.append(row)
    fmt = {c: ("0.000" if c.endswith("- p") else "0.00") for c in rows[0] if c != "Voorspeller"}
    add_sheet("1e_Regressie_tevredenheid", "Regressie: wat voorspelt tevredenheid met het leven?", VR, pd.DataFrame(rows),
              "Gewogen lineaire regressie (WLS), gestandaardiseerde coëfficiënten (beta): hoe groter, hoe sterker het verband. "
              "Alle modellen gecorrigeerd voor leeftijd, geslacht, opleiding en inkomen. Zelfde respondenten in alle modellen. "
              "p < 0,05 = statistisch significant. Toename R² van model 1 naar 2 = extra verklaring door toekomstverwachting.",
              fmt=fmt, visual="Tekst/kernboodschap; evt. staaf R² per model")

    # 1f per leeftijdsgroep (jongeren!)
    rows = []
    for lg in LEEFTIJD_LABELS + [TOTAAL]:
        m = (df["Leeftijdsgroep"] == lg) if lg != TOTAAL else df["Leeftijdsgroep"].notna()
        r, _ = wcorr(df.loc[m, "_tev_leven"], df.loc[m, "_toe_leven"], w[m])
        rows.append({"Leeftijdsgroep": lg,
                     "Gem. tevredenheid leven (1-10)": wmean(df.loc[m, "_tev_leven"], w[m]),
                     f"% tevreden ({TEVREDEN_GRENS}+)": wpct(df.loc[m, "_tevreden"], w[m]),
                     "% pessimistisch over leven": wpct(df.loc[m, "_pessimistisch"], w[m]),
                     "% optimistisch over leven": wpct(df.loc[m, "_optimistisch"], w[m]),
                     "% veel invloed op leven": wpct(df.loc[m, "_veel_invloed"], w[m]),
                     **{f"% {g}": wpct((df.loc[m, "_groep"] == g).astype(float).where(df.loc[m, "_groep"].notna()), w[m])
                        for g in GROEPEN},
                     "r tevredenheid x toekomst": r,
                     "n": int(m.sum())})
    add_sheet("1f_Per_leeftijdsgroep", "Tevredenheid, toekomstperspectief en grip per leeftijdsgroep", VR,
              pd.DataFrame(rows),
              "Let op: dit is één meetmoment; een trend ('steeds ontevredener') is met deze data niet te toetsen.",
              fmt={"Gem. tevredenheid leven (1-10)": "0.0", "r tevredenheid x toekomst": "0.00"},
              visual="Gegroepeerde staaf per leeftijdsgroep")


# ---------------------------------------------------------------- Subvraag 2
def tabellen_vraag2(df, V):
    w = df[W]
    VR = "2. Welke groepen zijn er (bijv. ontevreden & pessimistisch), wie zitten erin en waarover zijn ze pessimistisch?"

    def gw(mask_fn_series):
        # gewogen % binnen groep van een 0/1-reeks
        return per_groep(df, lambda m: wpct(mask_fn_series[m], w[m]))

    # 2a groepsoverzicht
    rows = [n_row(df, Kenmerk="n (ongewogen)"),
            {"Kenmerk": "% van de bevolking", **per_groep(df, lambda m: 100 * w[m].sum() / w[df["_groep"].notna()].sum())},
            {"Kenmerk": "Gem. tevredenheid leven (1-10)", **per_groep(df, lambda m: wmean(df.loc[m, "_tev_leven"], w[m]))},
            {"Kenmerk": "Gem. toekomstverwachting leven (1=veel slechter, 5=veel beter)", **per_groep(df, lambda m: wmean(df.loc[m, "_toe_leven"], w[m]))},
            {"Kenmerk": "% veel invloed op eigen leven", **gw(df["_veel_invloed"])},
            {"Kenmerk": "% geen invloed op eigen leven", **gw((df["_inv_leven"] == 0).astype(float).where(df["_inv_leven"].notna()))},
            {"Kenmerk": f"Gem. aantal domeinen onvoldoende (<{ONVOLDOENDE_GRENS})", **per_groep(df, lambda m: wmean(df.loc[m, "_n_onvoldoende"], w[m]))},
            {"Kenmerk": "Gem. aantal financiële onzekerheden", **per_groep(df, lambda m: wmean(df.loc[m, "_n_onz_fin"], w[m]))},
            {"Kenmerk": "Gem. leeftijd", **per_groep(df, lambda m: wmean(df.loc[m, "_leeftijd"], w[m]))}]
    if V["grip_stelling"]:
        s = V["stel"][V["grip_stelling"]]
        rows.append({"Kenmerk": "% (helemaal) eens: 'ik heb veel invloed op hoe mijn leven verloopt'",
                     **gw((s >= 4).astype(float).where(s.notna()))})
    add_sheet("2a_Groepen_overzicht", "De vier groepen: tevredenheid nu x verwachting voor de toekomst", VR,
              pd.DataFrame(rows),
              f"Tevreden = cijfer {TEVREDEN_GRENS}+ voor 'leven in het algemeen'. Pessimistisch = verwacht dat het leven over 5 jaar "
              "(een beetje of veel) slechter is; hoopvol/stabiel = verwacht hetzelfde of beter. 'Weet ik niet' = niet ingedeeld.",
              fmt={g: "0.0" for g in groep_cols()}, visual="Kwadrant/2x2-figuur met groepsgroottes")

    # 2b pessimisme per domein per groep
    rows = []
    for d in V["domeinen_excl_leven"]:  # 'leven' zelf is per definitie 0/100%
        if d not in V["toe"].columns:
            continue
        s = V["toe"][d]
        rows.append({"Domein": d, **gw((s < 3).astype(float).where(s.notna()))})
    tb = pd.DataFrame(rows)
    tb["Verschil ontevreden&pess. vs totaal (%-punt)"] = tb[GROEPEN[3]] - tb[TOTAAL]
    tb = tb.sort_values(GROEPEN[3], ascending=False)
    add_sheet("2b_Pessimisme_per_domein", "Op welke vlakken zijn groepen pessimistisch? (% verwacht 'slechter' over 5 jaar)",
              VR, tb, "Gewogen % dat verwacht dat het op dit domein (een beetje of veel) slechter gaat. "
                      "Gesorteerd op de groep 'Ontevreden & pessimistisch'.",
              visual="Gegroepeerde/dot-plot per domein")

    # 2c tevredenheid per domein per groep
    rows = []
    for d in V["domeinen_excl_leven"]:
        s = V["tev"][d]
        rows.append({"Domein": d, "Maat": "Gem. rapportcijfer",
                     **per_groep(df, lambda m, s=s: wmean(s[m], w[m]))})
        rows.append({"Domein": d, "Maat": f"% onvoldoende (<{ONVOLDOENDE_GRENS})",
                     **gw((s < ONVOLDOENDE_GRENS).astype(float).where(s.notna()))})
    add_sheet("2c_Tevredenheid_domein_groep", "Tevredenheid per levensdomein, per groep", VR, pd.DataFrame(rows),
              "Gewogen. Twee maten per domein: gemiddeld cijfer en % dat een onvoldoende geeft.",
              visual="Heatmap of gegroepeerde staaf")

    # 2d invloed per domein per groep
    rows = []
    for d in V["inv"].columns:
        s = V["inv"][d]
        rows.append({"Domein": d, "Maat": "% veel invloed", **gw((s == 2).astype(float).where(s.notna()))})
        rows.append({"Domein": d, "Maat": "% geen invloed", **gw((s == 0).astype(float).where(s.notna()))})
    for c in V["stel"].columns:
        s = V["stel"][c]
        rows.append({"Domein": f"Stelling: {c}", "Maat": "% (helemaal) eens", **gw((s >= 4).astype(float).where(s.notna()))})
    add_sheet("2d_Invloed_per_groep", "Ervaren invloed (grip) per domein, per groep", VR, pd.DataFrame(rows),
              "Gewogen. 'Weet ik niet' niet meegeteld.", visual="Gegroepeerde staaf")

    # 2e profiel: achtergrond
    def profiel(kenmerken, binblocks):
        rows = [n_row(df, Kenmerk="n (ongewogen)", Categorie="")]
        for naam, col in kenmerken:
            for cat in ordered_cats(df, col):
                ind = (df[col] == cat).astype(float).where(df[col].notna())
                rows.append({"Kenmerk": naam, "Categorie": cat, **gw(ind)})
        for naam, block in binblocks:
            for c in block.columns:
                rows.append({"Kenmerk": naam, "Categorie": c, **gw(block[c])})
        tb = pd.DataFrame(rows)
        tb["Index ontevreden&pess. (totaal=100)"] = 100 * tb[GROEPEN[3]] / tb[TOTAAL]
        tb.loc[0, "Index ontevreden&pess. (totaal=100)"] = np.nan
        return tb

    add_sheet("2e_Profiel_achtergrond", "Wie zitten er in elke groep? Achtergrondkenmerken", VR,
              profiel(V["achtergrond"], [("Dagbesteding (meerdere mogelijk)", V["dagbesteding"])]),
              "Kolom-percentages: verdeling van elk kenmerk BINNEN de groep (telt per kenmerk op tot 100%, behalve dagbesteding). "
              "Index > 100 = oververtegenwoordigd in 'Ontevreden & pessimistisch' t.o.v. totaal (bijv. 150 = 1,5x zo vaak).",
              fmt={"Index ontevreden&pess. (totaal=100)": "0"}, visual="Profielkaart per groep / gegroepeerde staaf")

    add_sheet("2f_Profiel_verleden", "Wie zitten er in elke groep? Vroegere en recente (financiële) situatie", VR,
              profiel(V["verleden"], []),
              "Kolom-percentages binnen de groep. Index > 100 = oververtegenwoordigd in 'Ontevreden & pessimistisch'.",
              fmt={"Index ontevreden&pess. (totaal=100)": "0"}, visual="Gegroepeerde staaf")

    # 2g onzekerheden & netwerk per groep
    rows = []
    for naam, block in [("Financiële/bestaansonzekerheid", V["onz_fin"]), ("Overige onzekerheid", V["onz_ovr"]),
                        ("Iemand die kan helpen met", V["net"]), ("Weet waar hulp te vinden bij", V["hulp"])]:
        for c in block.columns:
            rows.append({"Thema": naam, "Item": c, **gw(block[c])})
    tb = pd.DataFrame(rows)
    if len(tb):
        tb["Index ontevreden&pess. (totaal=100)"] = 100 * tb[GROEPEN[3]] / tb[TOTAAL]
    add_sheet("2g_Onzekerheid_netwerk_groep", "Onzekerheden, sociaal netwerk en hulp per groep", VR, tb,
              "Gewogen % 'van toepassing' / 'ja'. Bij netwerk en hulp telt 'weet ik niet' als 'nee' (geen zekere steun).",
              fmt={"Index ontevreden&pess. (totaal=100)": "0"}, visual="Gegroepeerde staaf")


# ---------------------------------------------------------------- Subvraag 3 & 4
def tabellen_vraag3_4(df, V):
    w = df[W]
    VR3 = "3. Is de kans op pessimisme groter als basisdingen (domeinen) niet op orde zijn?"
    VR4 = "4. Wat zijn 'basisdingen'? Welke domeinen wegen het zwaarst?"
    onv = V["onv"]

    # 3a pessimisme naar aantal onvoldoendes
    nb = df["_n_onvoldoende"].clip(upper=4)
    labels = {0: "0", 1: "1", 2: "2", 3: "3", 4: "4 of meer"}
    rows = []
    for k in range(5):
        m = nb == k
        rows.append({"Aantal domeinen met onvoldoende": labels[k],
                     "% van de bevolking": wpct(m.astype(float).where(nb.notna()), w),
                     "% pessimistisch over leven": wpct(df.loc[m, "_pessimistisch"], w[m]),
                     f"% tevreden met leven ({TEVREDEN_GRENS}+)": wpct(df.loc[m, "_tevreden"], w[m]),
                     "Gem. tevredenheid leven": wmean(df.loc[m, "_tev_leven"], w[m]),
                     "% Ontevreden & pessimistisch": wpct((df.loc[m, "_groep"] == GROEPEN[3]).astype(float).where(df.loc[m, "_groep"].notna()), w[m]),
                     "n": int(m.sum())})
    add_sheet("3a_Pessimisme_x_aantal_domeinen", "Hoe meer domeinen niet op orde, hoe pessimistischer?", VR3,
              pd.DataFrame(rows),
              f"Aantal van de {onv.shape[1]} levensdomeinen (excl. 'leven in het algemeen') waarop iemand een onvoldoende "
              f"(cijfer <{ONVOLDOENDE_GRENS}) geeft.",
              fmt={"Gem. tevredenheid leven": "0.0"}, visual="Lijn/staaf: oplopend % pessimistisch")

    # 3b per domein: kans op pessimisme bij onvoldoende vs voldoende
    C = controles(df, V)
    Cs = C.copy()
    Cs["Leeftijd"] = z(Cs["Leeftijd"])
    rows = []
    for d in onv.columns:
        m1, m0 = onv[d] == 1, onv[d] == 0
        p1, p0 = wpct(df.loc[m1, "_pessimistisch"], w[m1]), wpct(df.loc[m0, "_pessimistisch"], w[m0])
        try:
            mdl, n = fit_logit(df["_pessimistisch"], pd.concat([onv[[d]], Cs], axis=1), w)
            orr, (lo, hi), pv = np.exp(mdl.params[d]), np.exp(mdl.conf_int().loc[d]), mdl.pvalues[d]
        except Exception as e:  # noqa: BLE001
            log(f"  logit {d} mislukt: {e}")
            orr = lo = hi = pv = np.nan
        rows.append({"Domein": d,
                     f"% met onvoldoende (<{ONVOLDOENDE_GRENS})": wpct(onv[d], w),
                     "% pessimistisch als domein onvoldoende": p1,
                     "% pessimistisch als domein voldoende": p0,
                     "Verschil (%-punt)": p1 - p0,
                     "Relatieve kans (x zo vaak)": p1 / p0 if p0 else np.nan,
                     "Odds ratio (gecorrigeerd)": orr, "OR 95%-BI onder": lo, "OR 95%-BI boven": hi, "p": pv,
                     "n onvoldoende": int(m1.sum())})
    tb = pd.DataFrame(rows).sort_values("Verschil (%-punt)", ascending=False)
    add_sheet("3b_Pessimisme_x_domein", "Kans op pessimisme over het leven als een domein niet op orde is", VR3, tb,
              "Pessimistisch = verwacht dat het leven over 5 jaar slechter is. Odds ratio uit gewogen logistische regressie, "
              "per domein apart, gecorrigeerd voor leeftijd, geslacht, opleiding en inkomen (OR > 1 = grotere kans).",
              fmt={"Relatieve kans (x zo vaak)": "0.0", "Odds ratio (gecorrigeerd)": "0.00", "OR 95%-BI onder": "0.00",
                   "OR 95%-BI boven": "0.00", "p": "0.000"},
              visual="Gesorteerde staaf (verschil in %-punt)")

    # 4 basisdingen: welke domeinen wegen het zwaarst (multivariaat)
    doms = V["domeinen_excl_leven"]
    Xd = pd.DataFrame({d: z(V["tev"][d]) for d in doms})
    try:
        m_tev, n_tev = fit_wls(z(df["_tev_leven"]), pd.concat([Xd, Cs], axis=1), w)
    except Exception as e:  # noqa: BLE001
        log(f"  WLS basisdingen mislukt: {e}")
        m_tev, n_tev = None, 0
    try:
        m_pes, n_pes = fit_logit(df["_pessimistisch"], pd.concat([onv, Cs], axis=1), w)
    except Exception as e:  # noqa: BLE001
        log(f"  logit basisdingen mislukt: {e}")
        m_pes, n_pes = None, 0
    rows = []
    for d in doms:
        rows.append({"Domein": d,
                     "Beta op tevredenheid leven (samen)": m_tev.params.get(d, np.nan) if m_tev is not None else np.nan,
                     "p (tevredenheid)": m_tev.pvalues.get(d, np.nan) if m_tev is not None else np.nan,
                     "Odds ratio pessimisme bij onvoldoende (samen)": np.exp(m_pes.params.get(d, np.nan)) if m_pes is not None else np.nan,
                     "p (pessimisme)": m_pes.pvalues.get(d, np.nan) if m_pes is not None else np.nan,
                     f"% met onvoldoende (<{ONVOLDOENDE_GRENS})": wpct(onv[d], w),
                     "Gem. rapportcijfer": wmean(V["tev"][d], w)})
    tb = pd.DataFrame(rows)
    tb["Rang tevredenheid"] = tb["Beta op tevredenheid leven (samen)"].rank(ascending=False)
    tb["Rang pessimisme"] = tb["Odds ratio pessimisme bij onvoldoende (samen)"].rank(ascending=False)
    tb["Gemiddelde rang"] = tb[["Rang tevredenheid", "Rang pessimisme"]].mean(axis=1)
    tb = tb.sort_values("Gemiddelde rang")
    add_sheet("4_Basisdingen", "Welke levensdomeinen zijn 'basisdingen'? (alle domeinen tegelijk in één model)", VR4, tb,
              "Alle domeinen tegelijk in één model, zodat overlap tussen domeinen is weggefilterd. "
              f"Beta: gestandaardiseerd effect van domein-tevredenheid op tevredenheid met het leven (WLS, n={n_tev}). "
              f"Odds ratio: kans op pessimisme over het leven bij een onvoldoende op dat domein (logistisch, n={n_pes}). "
              "Gecorrigeerd voor leeftijd, geslacht, opleiding, inkomen. Laagste gemiddelde rang = meest 'basaal'.",
              fmt={"Beta op tevredenheid leven (samen)": "0.00", "p (tevredenheid)": "0.000",
                   "Odds ratio pessimisme bij onvoldoende (samen)": "0.00", "p (pessimisme)": "0.000",
                   "Gem. rapportcijfer": "0.0", "Rang tevredenheid": "0", "Rang pessimisme": "0", "Gemiddelde rang": "0.0"},
              visual="Top-5 lijst / gesorteerde staaf")


# ---------------------------------------------------------------- Subvraag 5
def tabellen_vraag5(df, V):
    w = df[W]
    VR = "5. Is zekerheid NU anders dan zekerheid over de TOEKOMST?"

    # 5a per domein: nu vs toekomst
    rows = []
    for d in V["domeinen"]:
        s, t = V["tev"][d], V["toe"].get(d)
        row = {"Domein": d,
               "Gem. rapportcijfer nu": wmean(s, w),
               f"% tevreden nu ({TEVREDEN_GRENS}+)": wpct((s >= TEVREDEN_GRENS).astype(float).where(s.notna()), w),
               f"% onvoldoende nu (<{ONVOLDOENDE_GRENS})": wpct((s < ONVOLDOENDE_GRENS).astype(float).where(s.notna()), w)}
        if t is not None:
            row.update({"% verwacht beter": wpct((t > 3).astype(float).where(t.notna()), w),
                        "% verwacht hetzelfde": wpct((t == 3).astype(float).where(t.notna()), w),
                        "% verwacht slechter": wpct((t < 3).astype(float).where(t.notna()), w)})
            row["Saldo toekomst (beter - slechter)"] = row["% verwacht beter"] - row["% verwacht slechter"]
            # zeker nu, onzeker later: tevreden nu maar verwacht slechter
            m = s >= TEVREDEN_GRENS
            row["% van tevredenen nu dat slechter verwacht"] = wpct((t[m] < 3).astype(float).where(t[m].notna()), w[m])
        rows.append(row)
    add_sheet("5a_Nu_vs_toekomst_per_domein", "Hoe staat het er nu voor en hoe kijkt men vooruit, per domein", VR,
              pd.DataFrame(rows),
              "Gewogen. Saldo > 0 = meer mensen verwachten verbetering dan verslechtering. "
              "Laatste kolom: 'nu zeker, later onzeker' - aandeel van wie nu tevreden is dat toch verslechtering verwacht.",
              fmt={"Gem. rapportcijfer nu": "0.0"}, visual="Dumbbell: % tevreden nu vs % verwacht slechter")

    # 5b onzekerheidsitems: verband met nu vs toekomst
    C = controles(df, V)
    Cs = C.copy()
    Cs["Leeftijd"] = z(Cs["Leeftijd"])
    items = pd.concat([V["onz_fin"], V["onz_ovr"]], axis=1)
    thema = {c: "Financieel/bestaan" for c in V["onz_fin"].columns}
    thema.update({c: "Overig" for c in V["onz_ovr"].columns})
    y_nu, y_toe = z(df["_tev_leven"]), z(df["_toe_leven"])
    m_nu = m_toe = None
    if items.shape[1]:
        try:
            m_nu, n_nu = fit_wls(y_nu, pd.concat([items, Cs], axis=1), w)
            m_toe, n_toe = fit_wls(y_toe, pd.concat([items, Cs], axis=1), w)
        except Exception as e:  # noqa: BLE001
            log(f"  WLS onzekerheid mislukt: {e}")
    rows = []
    for c in items.columns:
        m1, m0 = items[c] == 1, items[c] == 0
        b_nu = m_nu.params.get(c, np.nan) if m_nu is not None else np.nan
        b_toe = m_toe.params.get(c, np.nan) if m_toe is not None else np.nan
        rows.append({"Thema": thema[c], "Onzekerheid": c,
                     "% van toepassing": wpct(items[c], w),
                     "Gem. tevredenheid leven - wel": wmean(df.loc[m1, "_tev_leven"], w[m1]),
                     "Gem. tevredenheid leven - niet": wmean(df.loc[m0, "_tev_leven"], w[m0]),
                     "% pessimistisch - wel": wpct(df.loc[m1, "_pessimistisch"], w[m1]),
                     "% pessimistisch - niet": wpct(df.loc[m0, "_pessimistisch"], w[m0]),
                     "Effect op tevredenheid NU (SD)": b_nu,
                     "p (nu)": m_nu.pvalues.get(c, np.nan) if m_nu is not None else np.nan,
                     "Effect op verwachting TOEKOMST (SD)": b_toe,
                     "p (toekomst)": m_toe.pvalues.get(c, np.nan) if m_toe is not None else np.nan,
                     "Sterkst verbonden met": ("nu" if abs(b_nu) > abs(b_toe) else "toekomst") if pd.notna(b_nu) and pd.notna(b_toe) else "",
                     "n wel": int(m1.sum())})
    tb = pd.DataFrame(rows)
    add_sheet("5b_Onzekerheid_nu_vs_toekomst", "Hangen onzekerheden meer samen met tevredenheid nu of met verwachting voor de toekomst?",
              VR, tb,
              "Linker kolommen: gewogen beschrijvend (wel/niet van toepassing). Effect-kolommen: alle onzekerheden tegelijk in één "
              "gewogen regressie, gecorrigeerd voor leeftijd, geslacht, opleiding, inkomen; uitkomsten gestandaardiseerd zodat 'nu' "
              "(cijfer 1-10) en 'toekomst' (schaal 1-5) vergelijkbaar zijn. Negatief = lager cijfer / somberder verwachting.",
              fmt={"Gem. tevredenheid leven - wel": "0.0", "Gem. tevredenheid leven - niet": "0.0",
                   "Effect op tevredenheid NU (SD)": "0.00", "Effect op verwachting TOEKOMST (SD)": "0.00",
                   "p (nu)": "0.000", "p (toekomst)": "0.000"},
              visual="Twee staven per item (nu vs toekomst)")

    # 5c opeenstapeling onzekerheden
    nb = df["_n_onz_fin"].clip(upper=3)
    labels = {0: "0", 1: "1", 2: "2", 3: "3 of meer"}
    rows = []
    for k in range(4):
        m = nb == k
        rows.append({"Aantal financiële onzekerheden": labels[k],
                     "% van de bevolking": wpct(m.astype(float).where(nb.notna()), w),
                     "Gem. tevredenheid leven (nu)": wmean(df.loc[m, "_tev_leven"], w[m]),
                     f"% tevreden nu ({TEVREDEN_GRENS}+)": wpct(df.loc[m, "_tevreden"], w[m]),
                     "Gem. toekomstverwachting leven (1-5)": wmean(df.loc[m, "_toe_leven"], w[m]),
                     "% pessimistisch over leven": wpct(df.loc[m, "_pessimistisch"], w[m]),
                     "% veel invloed op leven": wpct(df.loc[m, "_veel_invloed"], w[m]),
                     "n": int(m.sum())})
    add_sheet("5c_Stapeling_onzekerheid", "Opeenstapeling van financiële onzekerheid: nu en toekomst", VR, pd.DataFrame(rows),
              "Aantal aangevinkte financiële/bestaansonzekerheden (schulden, geen vast inkomen, moeite met rondkomen, enz.).",
              fmt={"Gem. tevredenheid leven (nu)": "0.0", "Gem. toekomstverwachting leven (1-5)": "0.00"},
              visual="Twee lijnen: % tevreden nu en % pessimistisch")


# ---------------------------------------------------------------- Invloed / grip
def tabellen_invloed(df, V):
    w = df[W]
    VR = "6. Gevoel van invloed (grip/agency): hoe hangt het samen met tevredenheid en toekomstperspectief?"
    inv, tev, toe = V["inv"], V["tev"], V["toe"]
    lv = {0: "Geen invloed", 1: "Een beetje invloed", 2: "Veel invloed"}

    # 6a per domein: verdeling invloed + tevredenheid/pessimisme naar invloed
    rows = []
    for d in inv.columns:
        s = inv[d]
        raw = V["inv_raw"].get(d)
        row = {"Domein": d}
        for k, lab in lv.items():
            row[f"% {lab.lower()}"] = wpct((s == k).astype(float).where(s.notna()), w)
        if raw:  # 'weet ik niet' apart tonen (telt verder niet mee)
            r = norm_str(df[raw])
            row["% weet ik niet (van alle antwoorden)"] = wpct(r.str.contains("weet", na=False).astype(float).where(r.notna()), w)
        if d in tev.columns:
            for k in (0, 2):
                m = s == k
                row[f"Gem. cijfer domein - {lv[k].lower()}"] = wmean(tev.loc[m, d], w[m])
        if d in toe.columns:
            t = toe[d]
            for k in (0, 2):
                m = s == k
                row[f"% verwacht slechter - {lv[k].lower()}"] = wpct((t[m] < 3).astype(float).where(t[m].notna()), w[m])
        r1, _ = wcorr(s, tev[d], w) if d in tev.columns else (np.nan, 0)
        r2, _ = wcorr(s, toe[d], w) if d in toe.columns else (np.nan, 0)
        row["r invloed x tevredenheid"] = r1
        row["r invloed x toekomstverwachting"] = r2
        rows.append(row)
    add_sheet("6a_Invloed_per_domein", "Hoeveel invloed ervaren mensen per levensdomein, en maakt dat verschil?", VR,
              pd.DataFrame(rows),
              "Gewogen. Verdeling (geen/beetje/veel) exclusief 'weet ik niet'. Daarna: gemiddeld cijfer en % dat verslechtering "
              "verwacht, voor wie geen vs. veel invloed ervaart op dat domein. r = gewogen correlatie.",
              fmt={c: "0.0" for c in [f"Gem. cijfer domein - {lv[k].lower()}" for k in (0, 2)]} |
                  {"r invloed x tevredenheid": "0.00", "r invloed x toekomstverwachting": "0.00"},
              visual="100%-gestapelde staaf invloed per domein + verschilstaaf")

    # 6b 'vastzitten': onvoldoende én geen invloed
    rows = []
    for d in inv.columns:
        if d not in tev.columns:
            continue
        s, t = inv[d], tev[d]
        ok = s.notna() & t.notna()
        onv_ = t < ONVOLDOENDE_GRENS
        rows.append({"Domein": d,
                     "% onvoldoende én geen invloed ('vastzitten')": wpct((onv_ & (s == 0)).astype(float).where(ok), w),
                     "% onvoldoende, wel (enige) invloed": wpct((onv_ & (s >= 1)).astype(float).where(ok), w),
                     "% van wie onvoldoende geeft: geen invloed": wpct((s[onv_ & ok] == 0).astype(float), w[onv_ & ok]),
                     "n onvoldoende": int((onv_ & ok).sum())})
    tb = pd.DataFrame(rows).sort_values("% onvoldoende én geen invloed ('vastzitten')", ascending=False)
    add_sheet("6b_Vastzitten_per_domein", "Waar zitten mensen 'vast'? Ontevreden én geen gevoel van invloed", VR, tb,
              f"'Vastzitten' = onvoldoende (<{ONVOLDOENDE_GRENS}) voor een domein én ervaart er geen invloed op. "
              "Juist daar ligt een rol voor beleid: mensen kunnen het niet zelf oplossen.",
              visual="Gesorteerde staaf")

    # 6c buffer: helpt grip / netwerk als het tegenzit?
    C = controles(df, V)
    Cs = C.copy()
    Cs["Leeftijd"] = z(Cs["Leeftijd"])
    contexten = [
        ("1+ financiële onzekerheid", df["_n_onz_fin"] >= 1),
        ("Geen financiële onzekerheid", df["_n_onz_fin"] == 0),
        ("2+ domeinen onvoldoende", df["_n_onvoldoende"] >= 2),
        ("0-1 domeinen onvoldoende", df["_n_onvoldoende"] <= 1),
    ]
    beschermers = [("Ervaren invloed op leven", df["_inv_leven"], lv)]
    for key in ("Geld aan u te lenen", "luisterend oor"):
        c = next((c for c in V["net"].columns if key.lower() in c.lower()), None)
        if c:
            beschermers.append((f"Iemand die kan helpen: {c}", V["net"][c], {0: "Nee/weet niet", 1: "Ja"}))
    rows = []
    for ctx, cm in contexten:
        for bnaam, b, labels in beschermers:
            for k, lab in labels.items():
                m = cm & (b == k)
                rows.append({"Situatie": ctx, "Beschermende factor": bnaam, "Categorie": lab,
                             "% pessimistisch over leven": wpct(df.loc[m, "_pessimistisch"], w[m]),
                             "Gem. tevredenheid leven": wmean(df.loc[m, "_tev_leven"], w[m]),
                             f"% tevreden ({TEVREDEN_GRENS}+)": wpct(df.loc[m, "_tevreden"], w[m]),
                             "n": int(m.sum())})
    # toets of invloed sterker beschermt bij onzekerheid (interactie)
    toets = []
    for naam, risk in [("financiële onzekerheid", (df["_n_onz_fin"] >= 1).astype(float).where(df["_n_onz_fin"].notna())),
                       ("2+ domeinen onvoldoende", (df["_n_onvoldoende"] >= 2).astype(float).where(df["_n_onvoldoende"].notna()))]:
        try:
            X = pd.concat([risk.rename("risico"), df["_veel_invloed"].rename("veel_invloed"),
                           (risk * df["_veel_invloed"]).rename("interactie"), Cs], axis=1)
            mdl, _ = fit_logit(df["_pessimistisch"], X, w)
            toets.append(f"{naam}: interactie OR={np.exp(mdl.params['interactie']):.2f}, p={mdl.pvalues['interactie']:.3f}")
        except Exception as e:  # noqa: BLE001
            log(f"  buffer-toets {naam} mislukt: {e}")
    log("  Buffer-toetsen:", toets)
    add_sheet("6c_Buffer_invloed_netwerk", "Beschermt grip of een sociaal netwerk tegen pessimisme als het tegenzit?", VR,
              pd.DataFrame(rows),
              "Vergelijk binnen dezelfde situatie (bijv. 1+ financiële onzekerheid) wie wel/geen invloed of steun ervaart. "
              "Toets of veel invloed extra beschermt bij tegenslag (logistische regressie, interactie; OR<1 = extra bescherming): "
              + "; ".join(toets),
              fmt={"Gem. tevredenheid leven": "0.0"}, visual="Gegroepeerde staaf per situatie")

    # 6d regressie op toekomstverwachting: voegt grip iets toe bovenop tevredenheid nu?
    preds = {"Tevredenheid leven nu (1-10)": df["_tev_leven"],
             "Ervaren invloed op leven (0-2)": df["_inv_leven"]}
    for k, c in V["stel_kort"][:2]:
        preds[f"Stelling: {k} (1-5)"] = V["stel"][c]
    base = pd.concat([df["_toe_leven"].rename("_y")] + [s.rename(k) for k, s in preds.items()] + [C, w], axis=1).dropna()
    idx = base.index
    if len(idx) > 50:
        y = z(df.loc[idx, "_toe_leven"])
        Xp = pd.DataFrame({k: z(s.loc[idx]) for k, s in preds.items()})
        Cz = C.loc[idx].copy()
        Cz["Leeftijd"] = z(Cz["Leeftijd"])
        modellen = [("Model 1: achtergrond + tevredenheid nu", [list(preds)[0]]),
                    ("Model 2: + invloed/grip", list(preds))]
        res = {naam: fit_wls(y, pd.concat([Cz, Xp[p]], axis=1), w.loc[idx]) for naam, p in modellen}
        rows = []
        for k in list(preds) + ["Leeftijd"]:
            row = {"Voorspeller": k}
            for naam, _ in modellen:
                row[f"{naam} - beta"] = res[naam][0].params.get(k, np.nan)
                row[f"{naam} - p"] = res[naam][0].pvalues.get(k, np.nan)
            rows.append(row)
        rows.append({"Voorspeller": "R² (verklaarde variantie)", **{f"{n} - beta": res[n][0].rsquared for n, _ in modellen}})
        rows.append({"Voorspeller": "n", **{f"{n} - beta": res[n][1] for n, _ in modellen}})
        fmt = {c: ("0.000" if c.endswith("- p") else "0.00") for c in rows[0] if c != "Voorspeller"}
        add_sheet("6d_Regressie_toekomst", "Regressie: wat voorspelt de toekomstverwachting? Telt grip mee bovenop tevredenheid nu?",
                  VR, pd.DataFrame(rows),
                  "Gewogen lineaire regressie, gestandaardiseerde coëfficiënten (beta). Uitkomst: verwachting voor het leven over "
                  "5 jaar (1-5). Gecorrigeerd voor leeftijd, geslacht, opleiding, inkomen. Toename R² = extra verklaring door grip.",
                  fmt=fmt, visual="Tekst/kernboodschap")


# ---------------------------------------------------------------- Extra verhalen
def tabellen_extra(df, V):
    w = df[W]

    # 7a uitkomsten naar achtergrond en verleden (de 'kansenkaart')
    uitkomsten = {
        "Gem. tevredenheid leven (1-10)": lambda m: wmean(df.loc[m, "_tev_leven"], w[m]),
        f"% tevreden ({TEVREDEN_GRENS}+)": lambda m: wpct(df.loc[m, "_tevreden"], w[m]),
        "% pessimistisch over leven": lambda m: wpct(df.loc[m, "_pessimistisch"], w[m]),
        "% optimistisch over leven": lambda m: wpct(df.loc[m, "_optimistisch"], w[m]),
        "% veel invloed op leven": lambda m: wpct(df.loc[m, "_veel_invloed"], w[m]),
        "% Ontevreden & pessimistisch": lambda m: wpct((df.loc[m, "_groep"] == GROEPEN[3]).astype(float)
                                                       .where(df.loc[m, "_groep"].notna()), w[m]),
    }
    for k, c in V["stel_kort"]:
        s = V["stel"][c]
        uitkomsten[f"% eens: {k}"] = lambda m, s=s: wpct((s[m] >= 4).astype(float).where(s[m].notna()), w[m])
    rows = [{"Kenmerk": "Totaal", "Categorie": "", **{k: f(df["_tev_leven"].notna()) for k, f in uitkomsten.items()},
             "n": int(df["_tev_leven"].notna().sum())}]
    for naam, col in V["achtergrond"] + V["verleden"]:
        for cat in ordered_cats(df, col):
            m = df[col] == cat
            rows.append({"Kenmerk": naam, "Categorie": cat, **{k: f(m) for k, f in uitkomsten.items()}, "n": int(m.sum())})
    for c in V["dagbesteding"].columns:
        m = V["dagbesteding"][c] == 1
        rows.append({"Kenmerk": "Dagbesteding", "Categorie": c, **{k: f(m) for k, f in uitkomsten.items()}, "n": int(m.sum())})
    add_sheet("7a_Uitkomsten_naar_kenmerken", "Tevredenheid, toekomstblik en grip naar achtergrond en verleden",
              "Extra: welke groepen scoren laag? (o.a. kansengelijkheid: jeugd -> nu)", pd.DataFrame(rows),
              "Gewogen. Rij = groep mensen, kolom = uitkomst. Vergelijk met de rij 'Totaal'. Let op kleine n.",
              fmt={"Gem. tevredenheid leven (1-10)": "0.0"}, visual="Heatmap of 'kansenkaart' (rijen = groepen)")

    # 7b/7c redenen achter pessimisme en optimisme
    themas = [("Lichamelijke gezondheid", "lichamelijke gezondheid"), ("Woning", "woonsituatie"),
              ("Sociale contacten", "sociale contacten")]

    def redenen(richting):
        neg = richting == "neg"
        rows = []
        for dom, key in themas:
            cols = [c for c in find_cols(df, "u denkt dat uw " + key, "redenen")
                    if (("slechter" in clean_label(c).lower()) == neg)
                    and ("weet ik niet" not in subject(c).lower())]
            if not cols or dom not in V["toe"].columns:
                continue
            blk = binary_block(df, cols)
            opties = {}
            for c in cols:  # twee vragenlijstversies samenvoegen op antwoordoptie
                opt = re.sub(r"[\s:…\.]+$", "", subject(c))
                opties.setdefault(opt, []).append(c)
            sel = pd.DataFrame({o: blk[cs].max(axis=1) for o, cs in opties.items()})
            t = V["toe"][dom]
            doel = (t < 3) if neg else (t > 3)
            basis = doel & (sel.fillna(0).sum(axis=1) > 0)   # verwacht slechter/beter én noemde een reden
            subsets = [("Iedereen", basis)] + [(g, basis & (df["_groep"] == g)) for g in GROEPEN] + \
                      [(lg, basis & (df["Leeftijdsgroep"] == lg)) for lg in (LEEFTIJD_LABELS[0], LEEFTIJD_LABELS[-1])]
            for o in sel.columns:
                row = {"Domein": dom, "Reden": o}
                for naam, m in subsets:
                    row[naam] = wpct(sel.loc[m, o], w[m])
                rows.append(row)
            rows.append({"Domein": dom, "Reden": "n (ongewogen)", **{naam: int(m.sum()) for naam, m in subsets}})
        return pd.DataFrame(rows)

    for richting, naam, titel, uitleg in [
        ("neg", "7b_Redenen_pessimisme", "Waarom verwachten mensen dat het slechter gaat? (gezondheid, wonen, sociaal)",
         "Onder wie verwacht dat het domein slechter wordt én een reden aangaf. Meerdere redenen mogelijk (telt niet op tot 100%). "
         "Twee vragenlijstversies samengevoegd."),
        ("pos", "7c_Redenen_optimisme", "Waarom verwachten mensen dat het beter gaat? (gezondheid, wonen, sociaal)",
         "Onder wie verwacht dat het domein beter wordt én een reden aangaf. Meerdere redenen mogelijk."),
    ]:
        tb = redenen(richting)
        if len(tb):
            add_sheet(naam, titel, "Extra: achtergrond bij subvraag 2 (waarover en waarom pessimistisch)", tb, uitleg,
                      visual="Horizontale staaf per domein")
        else:
            log(f"  {naam}: geen redenkolommen gevonden")


def export_open_antwoorden(df, path):
    """Open antwoorden voor citaten. INTERN: niet naar de vormgever zonder anonimiteitscheck."""
    col = first_col(df, "Zijn er andere onderwerpen waarover u zorgen voor de toekomst heeft")
    if not col:
        log("  Open vraag niet gevonden")
        return
    s = df[col].astype("string").str.strip()
    m = s.notna() & (s.str.len() > 2) & ~s.str.lower().isin(["nee", "geen", "nvt", "n.v.t.", "-", "nee.", "geen."])
    out = pd.DataFrame({"Groep": df.loc[m, "_groep"], "Leeftijdsgroep": df.loc[m, "Leeftijdsgroep"],
                        "Tevredenheid leven": df.loc[m, "_tev_leven"],
                        "Verwachting leven (1-5)": df.loc[m, "_toe_leven"], "Antwoord": s[m]})
    out["_o"] = out["Groep"].map({g: i for i, g in enumerate(GROEPEN[::-1])}).fillna(9)
    out.sort_values(["_o", "Leeftijdsgroep"]).drop(columns="_o").to_excel(path, index=False)
    log(f"Open antwoorden (INTERN): {m.sum()} antwoorden -> {path}")


# ════════════════════════════════════════════════════════════════════════════
# 5. EXCEL SCHRIJVEN
# ════════════════════════════════════════════════════════════════════════════
def write_excel(path, n_resp):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    inhoud = pd.DataFrame([{"Tabblad": s["name"], "Titel": s["titel"], "Subvraag": s["vraag"],
                            "Suggestie visual": s["visual"]} for s in SHEETS])
    defin = [
        ("Bron", "Werkbestand BW en kansengelijkheid (respondenten met Reactiestatus ≠ 'Niet geantwoord')"),
        ("Aantal respondenten", n_resp),
        ("Weging", "Alle % en gemiddelden gewogen met 'weegfactor' (provincie x leeftijd x opleiding). n = ongewogen."),
        ("Tevreden", f"Rapportcijfer {TEVREDEN_GRENS} of hoger voor 'Uw leven in het algemeen'"),
        ("Onvoldoende / niet op orde", f"Rapportcijfer lager dan {ONVOLDOENDE_GRENS} voor een domein"),
        ("Pessimistisch", "Verwacht dat het over 5 jaar (een beetje of veel) slechter gaat"),
        ("Hoopvol/stabiel", "Verwacht dat het over 5 jaar hetzelfde of beter gaat"),
        ("Veel invloed (grip)", "Antwoord 'Veel invloed' op 'Hoeveel invloed denkt u dat u heeft op ...'"),
        ("Let op", "Percentages op basis van minder dan ca. 50 respondenten (kolom n) voorzichtig interpreteren."),
    ]

    head_fill = PatternFill("solid", fgColor="DDE7F0")
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        inhoud.to_excel(xw, sheet_name="0_Inhoud", startrow=4, index=False)
        ws = xw.sheets["0_Inhoud"]
        ws["A1"] = "Spoor 1 - Tevredenheid met het leven, toekomstperspectief en grip"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = "Cijfers voor het informatieproduct. Elk tabblad = één tabel/figuur."
        r0 = 5 + len(inhoud) + 2
        ws.cell(r0, 1, "Definities").font = Font(bold=True, size=12)
        for i, (k, v) in enumerate(defin, start=1):
            ws.cell(r0 + i, 1, k).font = Font(bold=True)
            ws.cell(r0 + i, 2, v)
        for c in range(1, 5):
            ws.cell(5, c).font = Font(bold=True)
            ws.cell(5, c).fill = head_fill
        for col, width in zip("ABCD", [32, 80, 70, 40]):
            ws.column_dimensions[col].width = width

        for s in SHEETS:
            d = s["df"].copy()
            d.to_excel(xw, sheet_name=s["name"], startrow=4, index=False)
            ws = xw.sheets[s["name"]]
            ws["A1"] = s["titel"]
            ws["A1"].font = Font(bold=True, size=14)
            ws["A2"] = s["vraag"]
            ws["A2"].font = Font(italic=True)
            ws["A3"] = s["toelichting"]
            ws["A3"].alignment = Alignment(wrap_text=False)
            ws["A4"] = "Bron: Werkbestand BW en kansengelijkheid, gewogen. n = ongewogen aantal."
            ws["A4"].font = Font(size=9, color="666666")
            for j, col in enumerate(d.columns, start=1):
                h = ws.cell(5, j)
                h.font = Font(bold=True)
                h.fill = head_fill
                h.alignment = Alignment(wrap_text=True, vertical="top")
                fmt = s["fmt"].get(col)
                if fmt is None:
                    if col in ("n", "n wel", "n onvoldoende") or str(col).startswith("n "):
                        fmt = "0"
                    elif pd.api.types.is_float_dtype(d[col]):
                        fmt = "0.0"
                if fmt:
                    for i in range(len(d)):
                        ws.cell(6 + i, j).number_format = fmt
                maxlen = max([len(str(col)) // 2 + 2] + [len(str(v)) for v in d[col].head(60).tolist()] + [8])
                ws.column_dimensions[get_column_letter(j)].width = min(max(maxlen + 2, 10), 60)
            ws.row_dimensions[5].height = 45
            ws.freeze_panes = ws.cell(6, 2)
            # n-rijen (rij 6 met 'n (ongewogen)') als geheel getal tonen
            first = d.iloc[0, 0] if len(d) else None
            if isinstance(first, str) and first.startswith("n (ongewogen)"):
                for j in range(1, d.shape[1] + 1):
                    ws.cell(6, j).number_format = "0"
                    ws.cell(6, j).font = Font(italic=True, color="666666")
    log(f"\nExcel geschreven: {path}  ({len(SHEETS) + 1} tabbladen)")


# ════════════════════════════════════════════════════════════════════════════
# 6. MAIN
# ════════════════════════════════════════════════════════════════════════════
def run(df, out_xlsx=OUT_XLSX, out_log=OUT_LOG):
    SHEETS.clear()
    df = df.loc[:, ~df.columns.duplicated()].copy()
    V = build_variables(df)
    for fn in (tabellen_vraag1, tabellen_vraag2, tabellen_vraag3_4, tabellen_vraag5,
               tabellen_invloed, tabellen_extra):
        log(f"\n── {fn.__name__} ──")
        fn(df, V)
    write_excel(out_xlsx, len(df))
    export_open_antwoorden(df, Path(out_xlsx).with_name("spoor1_open_antwoorden_INTERN.xlsx"))
    Path(out_log).write_text("\n".join(_LOG), encoding="utf-8")
    print(f"Controle-log: {out_log}")


if __name__ == "__main__":
    path = sys.argv[1] if len(sys.argv) > 1 else DATA_PATH
    run(load_data(path))
